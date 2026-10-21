import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from app.config import settings
from app.db import db_session, initialize
from fastapi import HTTPException

from app.main import completed_plan_rows, completed_plans_workbook, export_completed_plans


class CompletedPlansExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.original_database_path = settings.database_path
        object.__setattr__(
            settings,
            "database_path",
            Path(self.temporary_directory.name) / "export-test.sqlite3",
        )
        initialize()

    def tearDown(self):
        object.__setattr__(settings, "database_path", self.original_database_path)
        self.temporary_directory.cleanup()

    def insert_concluded_plan(self):
        evaluation = {
            "ID": "E-1",
            "Fecha": "2026-10-03",
            "CveUnidad": "F0246",
            "Unidad": "VALDIVIA 7",
            "NomSupervisor": "SUPERVISOR PRUEBA",
            "NomJefeComercial": "JEFE PRUEBA",
            "Fuente": "Google",
            "TipoEvaluacion": "Mala",
            "Comentario": "Comentario de prueba",
        }
        payload = json.dumps(evaluation, ensure_ascii=False, sort_keys=True)
        payload_hash = hashlib.sha256(payload.encode("utf-8")).hexdigest()
        with db_session() as connection:
            run_id = connection.execute(
                """INSERT INTO evaluation_sync_runs
                   (trigger_type, status, week_start, week_end, started_at, completed_at)
                   VALUES ('test', 'success', '2026-09-28', '2026-10-04',
                           '2026-10-05T00:00:00Z', '2026-10-05T00:01:00Z')"""
            ).lastrowid
            connection.execute(
                """INSERT INTO evaluation_history
                   (attention_id, evaluation_date, payload_json, payload_hash,
                    first_seen_run_id, last_seen_run_id, created_at, updated_at)
                   VALUES ('E-1', '2026-10-03', ?, ?, ?, ?,
                           CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)""",
                (payload, payload_hash, run_id, run_id),
            )
            connection.execute(
                """INSERT INTO action_plans
                   (attention_id, status, plan_text, due_date, created_by, updated_by,
                    created_at, updated_at, completed_at)
                   VALUES ('E-1', 'concluido', 'Acción terminada', '2026-10-10',
                           'supervisor', 'supervisor', '2026-10-05 13:30:00',
                           '2026-10-08 15:00:00', '2026-10-08 15:00:00')"""
            )

    def test_export_uses_historical_concluded_plans_and_calculates_days(self):
        self.insert_concluded_plan()

        rows = completed_plan_rows({"role": "admin", "scope_name": None})

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Estado"], "Concluido")
        self.assertEqual(rows[0]["Días hasta crear el plan"], 2)
        self.assertEqual(rows[0]["Días de gestión del plan"], 3)
        self.assertEqual(rows[0]["Días desde evaluación hasta conclusión"], 5)

    def test_workbook_is_formatted_and_filterable(self):
        self.insert_concluded_plan()
        rows = completed_plan_rows({"role": "admin", "scope_name": None})

        workbook = load_workbook(completed_plans_workbook(rows))
        sheet = workbook["Planes concluidos"]

        self.assertTrue(sheet["A1"].value.startswith("EXPERIENCIA ATENCIÓN"))
        self.assertEqual(sheet["K5"].value, "Concluido")
        self.assertEqual(sheet["P5"].value, 2)
        self.assertEqual(sheet.auto_filter.ref, "A4:T5")
        self.assertEqual(sheet.freeze_panes, "A5")

    def test_source_text_is_not_executed_as_an_excel_formula(self):
        self.insert_concluded_plan()
        rows = completed_plan_rows({"role": "admin", "scope_name": None})
        rows[0]["Comentario de la evaluación"] = "=1+1"

        sheet = load_workbook(completed_plans_workbook(rows))["Planes concluidos"]

        self.assertEqual(sheet["J5"].data_type, "s")
        self.assertEqual(sheet["J5"].value, "'=1+1")

    def test_export_respects_the_logged_in_users_scope(self):
        self.insert_concluded_plan()

        visible = completed_plan_rows({
            "role": "supervisor_franquicias",
            "scope_name": "SUPERVISOR PRUEBA",
        })
        hidden = completed_plan_rows({
            "role": "supervisor_franquicias",
            "scope_name": "OTRO SUPERVISOR",
        })

        self.assertEqual(len(visible), 1)
        self.assertEqual(hidden, [])

    def test_export_reports_when_there_are_no_completed_plans(self):
        with self.assertRaises(HTTPException) as captured:
            export_completed_plans({"role": "admin", "scope_name": None})

        self.assertEqual(captured.exception.status_code, 404)
        self.assertEqual(
            captured.exception.detail,
            "No hay planes concluidos para exportar.",
        )


if __name__ == "__main__":
    unittest.main()
