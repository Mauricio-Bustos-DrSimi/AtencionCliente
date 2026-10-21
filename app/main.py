import csv
import io
import logging
import secrets
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import Depends, FastAPI, Header, HTTPException, Request, status
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from starlette.middleware.sessions import SessionMiddleware

from .config import settings
from .db import (
    PLAN_STATUSES,
    FIXED_ADMIN_USERNAME,
    ROLES,
    admin_password_is_ready,
    clear_login_failures,
    delete_user_account,
    delete_plan,
    get_latest_sync_run,
    get_plans,
    get_plan_history,
    get_stored_evaluations,
    get_user,
    initialize,
    list_users,
    login_is_locked,
    record_login_failure,
    save_plan,
    save_user,
    save_users_batch,
    update_user_account,
)
from .security import hash_password, verify_password


logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("atencion_cliente")
STATIC_DIR = Path(__file__).resolve().parent / "static"
IMAGE_DIR = Path(__file__).resolve().parent.parent / "image"

app = FastAPI(title="Experiencia Atención al Cliente", docs_url=None, redoc_url=None)
app.add_middleware(
    SessionMiddleware,
    secret_key=settings.secret_key,
    session_cookie="atencion_session",
    same_site="lax",
    https_only=settings.session_https_only,
    max_age=8 * 60 * 60,
)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/images", StaticFiles(directory=IMAGE_DIR), name="images")


@app.middleware("http")
async def prevent_stale_application_data(request: Request, call_next):
    response = await call_next(request)
    if request.url.path == "/" or request.url.path.startswith(("/api/", "/static/")):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; "
        "img-src 'self' data:; font-src 'self' data:; connect-src 'self'; "
        "object-src 'none'; base-uri 'none'; form-action 'self'; frame-ancestors 'none'"
    )
    return response


class LoginBody(BaseModel):
    username: str = Field(min_length=1, max_length=100)
    password: str = Field(min_length=1, max_length=200)


class PlanBody(BaseModel):
    status: str = Field(pattern="^(pendiente|en_subsanacion|concluido)$")
    plan_text: str = Field(default="", max_length=10000)
    due_date: Optional[date] = None
    completion_confirmed: bool = False


class UserBody(BaseModel):
    username: str = Field(min_length=3, max_length=100)
    full_name: str = Field(min_length=2, max_length=160)
    password: str = Field(min_length=12, max_length=200)
    role: str
    scope_name: Optional[str] = Field(default=None, max_length=160)


class UsersImportBody(BaseModel):
    filename: str = Field(min_length=1, max_length=255)
    content: str = Field(min_length=1, max_length=2_000_000)


class UserUpdateBody(BaseModel):
    full_name: str = Field(min_length=2, max_length=160)
    role: str
    scope_name: Optional[str] = Field(default=None, max_length=160)
    password: Optional[str] = Field(default=None, max_length=200)


@app.on_event("startup")
def startup() -> None:
    if len(settings.secret_key) < 32 or settings.secret_key == "replace-with-a-random-secret":
        raise RuntimeError("Configura APP_SECRET_KEY con al menos 32 caracteres aleatorios.")
    initialize()
    if not admin_password_is_ready():
        raise RuntimeError(
            "Configura la contraseña de admin con: python -m app.cli set-password --username admin"
        )


@app.get("/", include_in_schema=False)
def home():
    return FileResponse(
        STATIC_DIR / "index.html",
        headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"},
    )


@app.get("/api/health")
def health():
    return {"status": "ok", "service": "atencion-cliente"}


def current_user(request: Request):
    username = request.session.get("username")
    if not username:
        raise HTTPException(status_code=401, detail="Inicia sesión para continuar.")
    user = get_user(username)
    if (
        not user
        or request.session.get("auth_version") != user["auth_version"]
        or (user["role"] == "admin" and user["username"].casefold() != FIXED_ADMIN_USERNAME)
    ):
        request.session.clear()
        raise HTTPException(status_code=401, detail="La sesión ya no está activa.")
    return dict(user)


def require_csrf(request: Request, token: Optional[str]) -> None:
    expected = request.session.get("csrf_token")
    if not expected or not token or not secrets.compare_digest(expected, token):
        raise HTTPException(status_code=403, detail="Token de seguridad inválido.")


def normalized_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(character for character in text if not unicodedata.combining(character)).strip().casefold()


def parse_users_import(content: str) -> tuple[list[dict], list[dict]]:
    """Lee CSV o TXT delimitado y devuelve filas válidas y filas omitidas."""
    try:
        dialect = csv.Sniffer().sniff(content[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(content), dialect=dialect)
    if not reader.fieldnames:
        raise HTTPException(status_code=422, detail="El archivo debe incluir una fila de encabezados.")

    headers = {normalized_header(header): header for header in reader.fieldnames if header}
    aliases = {
        "username": ("usuario", "username", "user"),
        "role": ("rol", "role"),
        "password": ("contrasena", "password", "clave"),
        "full_name": ("nombre", "nombre completo", "full_name", "nombrecompleto"),
        "scope_name": ("alcance", "scope", "scope_name"),
    }
    selected = {
        field: next((headers[alias] for alias in names if alias in headers), None)
        for field, names in aliases.items()
    }
    required = ("username", "role", "password")
    missing = [field for field in required if not selected[field]]
    if missing:
        labels = {"username": "usuario", "role": "rol", "password": "contraseña"}
        raise HTTPException(
            status_code=422,
            detail=f"Faltan las columnas requeridas: {', '.join(labels[field] for field in missing)}.",
        )

    scoped_roles = {
        "jefe_comercial_ventas", "jefe_comercial_franquicias",
        "supervisor_ventas", "supervisor_franquicias", "gerencia",
    }
    users: list[dict] = []
    skipped: list[dict] = []
    for row in reader:
        line = reader.line_num
        if not any(str(value or "").strip() for value in row.values()):
            continue
        username = str(row.get(selected["username"]) or "").strip()
        role = str(row.get(selected["role"]) or "").strip()
        password = str(row.get(selected["password"]) or "")
        full_name = str(row.get(selected["full_name"]) or username).strip()
        scope_name = str(row.get(selected["scope_name"]) or "").strip()
        reason = None
        if not 3 <= len(username) <= 100:
            reason = "El usuario debe tener entre 3 y 100 caracteres."
        elif username.casefold() == FIXED_ADMIN_USERNAME:
            reason = "El usuario admin es fijo y no se puede importar."
        elif role == "admin":
            reason = "El único rol Administrador corresponde al usuario admin fijo."
        elif role not in ROLES:
            reason = "El rol no es válido."
        elif not 12 <= len(password) <= 200:
            reason = "La contraseña debe tener entre 12 y 200 caracteres."
        elif not 2 <= len(full_name) <= 160:
            reason = "El nombre debe tener entre 2 y 160 caracteres."
        elif role in scoped_roles and not scope_name:
            reason = "Este rol requiere la columna opcional alcance con un valor."
        if reason:
            skipped.append({"line": line, "reason": reason})
            continue
        try:
            password_hash = hash_password(password)
        except ValueError as exc:
            skipped.append({"line": line, "reason": str(exc)})
            continue
        users.append({
            "line": line,
            "username": username,
            "full_name": full_name,
            "password_hash": password_hash,
            "role": role,
            "scope_name": scope_name,
        })
    return users, skipped


def validate_user_assignment(username: str, role: str, scope_name: str, password: str | None = None) -> None:
    if role not in ROLES:
        raise HTTPException(status_code=422, detail="Rol no reconocido.")
    if username.casefold() == FIXED_ADMIN_USERNAME:
        if role != "admin":
            raise HTTPException(status_code=422, detail="El usuario admin debe conservar el rol Administrador.")
    elif role == "admin":
        raise HTTPException(status_code=422, detail="El único Administrador permitido es el usuario admin fijo.")
    scoped_roles = {
        "jefe_comercial_ventas", "jefe_comercial_franquicias",
        "supervisor_ventas", "supervisor_franquicias", "gerencia",
    }
    if role in scoped_roles and not scope_name:
        raise HTTPException(status_code=422, detail="Este rol requiere un alcance asignado.")
    if password is not None and not 12 <= len(password) <= 200:
        raise HTTPException(status_code=422, detail="La contraseña debe tener entre 12 y 200 caracteres.")


def same_name(value, scope) -> bool:
    return bool(value and scope and str(value).strip().casefold() == str(scope).strip().casefold())


def division_for(row: dict) -> str:
    code = str(row.get("CveUnidad") or "").strip().upper()
    if code.startswith("CL"):
        return "Sucursal"
    if code.startswith("F"):
        return "Franquicia"
    return ""


def visible_to_user(row: dict, user: dict) -> bool:
    role = user["role"]
    division = division_for(row)
    if role in {"admin", "gerente_general", "gerencia", "viewer"}:
        return True
    if role == "gerente_ventas":
        return division == "Sucursal"
    if role == "gerente_franquicias":
        return division == "Franquicia"
    if role == "jefe_comercial_ventas":
        return division == "Sucursal" and same_name(
            row.get("NomJefeComercial"), user.get("scope_name")
        )
    if role == "jefe_comercial_franquicias":
        return division == "Franquicia" and same_name(
            row.get("NomJefeComercial"), user.get("scope_name")
        )
    if role == "supervisor_ventas":
        return division == "Sucursal" and same_name(
            row.get("NomSupervisor"), user.get("scope_name")
        )
    if role == "supervisor_franquicias":
        return division == "Franquicia" and same_name(
            row.get("NomSupervisor"), user.get("scope_name")
        )
    return False


def load_visible_rows(user: dict, view: str = "historical") -> list[dict]:
    rows = get_stored_evaluations(view)
    return [row for row in rows if visible_to_user(row, user)]


def category(row: dict) -> str:
    value = str(row.get("TipoEvaluacion") or "").strip().casefold()
    if value in {"mala", "regular", "buena"}:
        return value.title()
    if str(row.get("Fuente") or "").upper() == "MS":
        try:
            score = float(row.get("PuntajeMS"))
        except (TypeError, ValueError):
            return "Sin clasificación"
        if 0 <= score <= 49:
            return "Mala"
        if 50 <= score <= 79:
            return "Regular"
        if 80 <= score <= 100:
            return "Buena"
        return "Sin clasificación"
    try:
        rating = int(row.get("Rating"))
    except (TypeError, ValueError):
        return "Sin clasificación"
    if rating in {1, 2}:
        return "Mala"
    if rating == 3:
        return "Regular"
    if rating in {4, 5}:
        return "Buena"
    return "Sin clasificación"


def parse_evaluation_date(value) -> date | None:
    try:
        return date.fromisoformat(str(value)[:10])
    except (TypeError, ValueError):
        return None


def summary_for(evaluations: list[dict]) -> dict:
    summary = {
        "evaluaciones": len(evaluations),
        "malas": sum(item["category"] == "Mala" for item in evaluations),
        "regulares": sum(item["category"] == "Regular" for item in evaluations),
        "buenas": sum(item["category"] == "Buena" for item in evaluations),
        "planes_pendientes": sum(
            item["plan_required"] and item["plan_status"] == "pendiente"
            for item in evaluations
        ),
        "franquicias": sum(division_for(item) == "Franquicia" for item in evaluations),
        "sucursales": sum(division_for(item) == "Sucursal" for item in evaluations),
    }
    for key, label in (("malas", "Mala"), ("regulares", "Regular"), ("buenas", "Buena")):
        summary[f"{key}_franquicias"] = sum(
            item["category"] == label and division_for(item) == "Franquicia"
            for item in evaluations
        )
        summary[f"{key}_sucursales"] = sum(
            item["category"] == label and division_for(item) == "Sucursal"
            for item in evaluations
        )
    summary["evaluaciones_franquicias"] = summary["franquicias"]
    summary["evaluaciones_sucursales"] = summary["sucursales"]
    return summary


def period_label(start: date, end: date) -> str:
    months = (
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
    )
    week = start.isocalendar().week
    start_day = str(start.day).zfill(2)
    end_day = str(end.day).zfill(2)
    if start.month == end.month:
        return f"Semana {week}: {start_day} al {end_day} de {months[end.month - 1]} del {end.year}"
    return (
        f"Semana {week}: {start_day} de {months[start.month - 1]} "
        f"al {end_day} de {months[end.month - 1]} del {end.year}"
    )


def local_datetime(value: object) -> datetime | None:
    """Convierte las fechas UTC de SQLite a la hora local de Santiago."""
    text = str(value or "").strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(ZoneInfo("America/Santiago"))


def completed_plan_rows(user: dict) -> list[dict]:
    """Reúne todos los planes concluidos visibles para el usuario."""
    plans = get_plans()
    rows = []
    for evaluation in load_visible_rows(user, "historical"):
        attention_id = str(evaluation.get("ID") or "").strip()
        plan = plans.get(attention_id)
        if not plan or plan["status"] != "concluido":
            continue
        evaluation_day = parse_evaluation_date(evaluation.get("Fecha"))
        created_at = local_datetime(plan.get("created_at"))
        completed_at = local_datetime(plan.get("completed_at"))
        days_to_create = (
            max((created_at.date() - evaluation_day).days, 0)
            if evaluation_day and created_at else None
        )
        management_days = (
            max((completed_at.date() - created_at.date()).days, 0)
            if created_at and completed_at else None
        )
        days_to_complete = (
            max((completed_at.date() - evaluation_day).days, 0)
            if evaluation_day and completed_at else None
        )
        rows.append({
            "ID evaluación": attention_id,
            "Código de unidad": evaluation.get("CveUnidad") or "",
            "Unidad": evaluation.get("Unidad") or "",
            "Tipo de local": division_for(evaluation),
            "Supervisor": evaluation.get("NomSupervisor") or "",
            "Jefe comercial": evaluation.get("NomJefeComercial") or "",
            "Fuente": evaluation.get("Fuente") or "",
            "Fecha de evaluación": evaluation_day,
            "Clasificación": category(evaluation),
            "Comentario de la evaluación": evaluation.get("Comentario") or "",
            "Estado": "Concluido",
            "Compromisos, comentarios y acciones": plan.get("plan_text") or "",
            "Fecha de plazo": parse_evaluation_date(plan.get("due_date")),
            "Creado por": plan.get("created_by") or "",
            "Fecha de creación del plan": created_at,
            "Días hasta crear el plan": days_to_create,
            "Concluido por": plan.get("updated_by") or "",
            "Fecha de conclusión": completed_at,
            "Días de gestión del plan": management_days,
            "Días desde evaluación hasta conclusión": days_to_complete,
        })
    return sorted(
        rows,
        key=lambda item: item["Fecha de conclusión"] or datetime.min.replace(tzinfo=timezone.utc),
        reverse=True,
    )


def completed_plans_workbook(rows: list[dict]) -> io.BytesIO:
    """Construye un Excel legible y filtrable con los planes concluidos."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Planes concluidos"
    headers = list(rows[0]) if rows else [
        "ID evaluación", "Código de unidad", "Unidad", "Tipo de local",
        "Supervisor", "Jefe comercial", "Fuente", "Fecha de evaluación",
        "Clasificación", "Comentario de la evaluación", "Estado",
        "Compromisos, comentarios y acciones", "Fecha de plazo", "Creado por",
        "Fecha de creación del plan", "Días hasta crear el plan", "Concluido por",
        "Fecha de conclusión", "Días de gestión del plan",
        "Días desde evaluación hasta conclusión",
    ]
    last_column = get_column_letter(len(headers))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = "EXPERIENCIA ATENCIÓN AL CLIENTE · PLANES DE ACCIÓN CONCLUIDOS"
    sheet["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="123B5D")
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = (
        f"Registros: {len(rows)} · Generado: "
        f"{datetime.now(ZoneInfo('America/Santiago')).strftime('%d-%m-%Y %H:%M')}"
    )
    sheet["A2"].font = Font(name="Aptos", size=10, italic=True, color="52677A")
    sheet["A2"].fill = PatternFill("solid", fgColor="EAF2F7")
    sheet["A2"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 22

    header_row = 4
    header_fill = PatternFill("solid", fgColor="2877A6")
    header_font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9E3EA"),
        right=Side(style="thin", color="D9E3EA"),
        top=Side(style="thin", color="D9E3EA"),
        bottom=Side(style="thin", color="D9E3EA"),
    )
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    sheet.row_dimensions[header_row].height = 34

    for row_number, item in enumerate(rows, header_row + 1):
        row_fill = PatternFill("solid", fgColor="F5F9FC" if row_number % 2 else "FFFFFF")
        for column, header in enumerate(headers, 1):
            value = item.get(header)
            if isinstance(value, datetime) and value.tzinfo:
                value = value.replace(tzinfo=None)
            if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
                value = "'" + value
            cell = sheet.cell(row_number, column, value)
            cell.fill = PatternFill("solid", fgColor="E5F6EC") if header == "Estado" else row_fill
            cell.font = Font(name="Aptos", size=10, bold=header == "Estado", color="217447" if header == "Estado" else "24364A")
            cell.alignment = Alignment(vertical="top", wrap_text=header in {
                "Comentario de la evaluación", "Compromisos, comentarios y acciones"
            })
            cell.border = thin_border
            if isinstance(value, datetime):
                cell.number_format = "dd-mm-yyyy hh:mm"
            elif isinstance(value, date):
                cell.number_format = "dd-mm-yyyy"
        sheet.row_dimensions[row_number].height = 42

    widths = {
        "ID evaluación": 23, "Código de unidad": 16, "Unidad": 28,
        "Tipo de local": 16, "Supervisor": 28, "Jefe comercial": 28,
        "Fuente": 14, "Fecha de evaluación": 18, "Clasificación": 15,
        "Comentario de la evaluación": 48, "Estado": 14,
        "Compromisos, comentarios y acciones": 55, "Fecha de plazo": 16,
        "Creado por": 18, "Fecha de creación del plan": 22,
        "Días hasta crear el plan": 20, "Concluido por": 18,
        "Fecha de conclusión": 21, "Días de gestión del plan": 22,
        "Días desde evaluación hasta conclusión": 28,
    }
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(header, 18)
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{last_column}{max(header_row + len(rows), header_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.add_filter_column(headers.index("Estado"), ["Concluido"])
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@app.post("/api/login")
def login(body: LoginBody, request: Request):
    username = body.username.strip()
    if login_is_locked(username):
        raise HTTPException(status_code=429, detail="Demasiados intentos. Intenta nuevamente en 15 minutos.")
    user = get_user(username)
    if user and user["role"] == "admin" and user["username"].casefold() != FIXED_ADMIN_USERNAME:
        record_login_failure(username)
        raise HTTPException(status_code=401, detail="Usuario o contraseña incorrectos.")
    if user and user["role"] == "admin" and not user["password_ready"]:
        raise HTTPException(status_code=503, detail="La cuenta administradora requiere una contraseña nueva.")
    if not user or not verify_password(body.password, user["password_hash"]):
        record_login_failure(username)
        raise HTTPException(status_code=401, detail="Usuario o contraseña incorrectos.")
    clear_login_failures(username)
    request.session.clear()
    request.session["username"] = user["username"]
    request.session["auth_version"] = user["auth_version"]
    request.session["csrf_token"] = secrets.token_urlsafe(32)
    return {
        "username": user["username"],
        "full_name": user["full_name"],
        "role": user["role"],
        "role_label": ROLES[user["role"]],
        "csrf_token": request.session["csrf_token"],
        "can_manage_plans": user["role"] != "viewer",
    }


@app.post("/api/logout")
def logout(
    request: Request,
    user=Depends(current_user),
    x_csrf_token: Optional[str] = Header(default=None),
):
    require_csrf(request, x_csrf_token)
    request.session.clear()
    return {"ok": True}


@app.get("/api/me")
def me(request: Request, user=Depends(current_user)):
    return {
        "username": user["username"],
        "full_name": user["full_name"],
        "role": user["role"],
        "role_label": ROLES[user["role"]],
        "scope_name": user["scope_name"],
        "csrf_token": request.session.get("csrf_token"),
        "can_manage_plans": user["role"] != "viewer",
    }


@app.get("/api/dashboard")
def dashboard(view: str = "weekly", user=Depends(current_user)):
    if view not in {"weekly", "historical"}:
        raise HTTPException(status_code=422, detail="Vista de evaluaciones no válida.")
    rows = load_visible_rows(user, view)
    plans = get_plans()
    evaluations = []
    for row in rows:
        item = dict(row)
        item["plan"] = plans.get(item["ID"])
        item["plan_status"] = item["plan"]["status"] if item["plan"] else "pendiente"
        item["category"] = category(item)
        item["plan_required"] = item["category"] == "Mala"
        evaluations.append(item)

    latest_sync = get_latest_sync_run()
    latest_success = get_latest_sync_run(successful_only=True)
    period = None
    if view == "weekly" and latest_success:
        period_start = date.fromisoformat(latest_success["week_start"])
        period_end = date.fromisoformat(latest_success["week_end"])
        period = {
            "start": period_start.isoformat(),
            "end": period_end.isoformat(),
            "label": period_label(period_start, period_end),
        }
    elif view == "historical":
        available_dates = sorted(
            evaluation_date
            for item in evaluations
            if (evaluation_date := parse_evaluation_date(item.get("Fecha")))
        )
        if available_dates:
            start, end = available_dates[0], available_dates[-1]
            period = {
                "start": start.isoformat(),
                "end": end.isoformat(),
                "label": f"Histórico acumulado: {start.strftime('%d-%m-%Y')} al {end.strftime('%d-%m-%Y')}",
            }
    return {
        "view": view,
        "summary": summary_for(evaluations),
        "evaluations": evaluations,
        "period": period,
        "last_sync": latest_sync,
    }


@app.get("/api/exports/completed-plans")
def export_completed_plans(user=Depends(current_user)):
    """Exporta todos los planes concluidos dentro del alcance del usuario."""
    rows = completed_plan_rows(user)
    if not rows:
        raise HTTPException(
            status_code=404,
            detail="No hay planes concluidos para exportar.",
        )
    filename = (
        "Planes_Accion_Concluidos_"
        f"{datetime.now(ZoneInfo('America/Santiago')).strftime('%Y%m%d_%H%M')}.xlsx"
    )
    return StreamingResponse(
        completed_plans_workbook(rows),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/attention/{attention_id}/history")
def action_history(attention_id: str, user=Depends(current_user)):
    rows = load_visible_rows(user, "historical")
    if not any(row["ID"] == attention_id for row in rows):
        raise HTTPException(status_code=404, detail="Evaluaci\u00f3n no encontrada.")
    return {"transactions": get_plan_history(attention_id)}


@app.put("/api/attention/{attention_id}/plan")
def update_plan(
    attention_id: str,
    body: PlanBody,
    request: Request,
    user=Depends(current_user),
    x_csrf_token: Optional[str] = Header(default=None),
):
    require_csrf(request, x_csrf_token)
    if user["role"] == "viewer":
        raise HTTPException(status_code=403, detail="El perfil Visualizador no puede crear ni modificar planes.")
    rows = load_visible_rows(user, "historical")
    attention = next((row for row in rows if row["ID"] == attention_id), None)
    if not attention:
        raise HTTPException(status_code=404, detail="Evaluación no encontrada.")
    if body.status not in PLAN_STATUSES:
        raise HTTPException(status_code=422, detail="Estado de plan no válido.")
    existing_plan = get_plans().get(attention_id)
    if (
        body.status == "concluido"
        and (not existing_plan or existing_plan["status"] != "concluido")
        and not body.completion_confirmed
    ):
        raise HTTPException(
            status_code=409,
            detail="Confirma expresamente la conclusión del plan de acción.",
        )

    plan = save_plan(
        attention_id=attention_id,
        status=body.status,
        plan_text=body.plan_text.strip(),
        due_date=body.due_date.isoformat() if body.due_date else None,
        username=user["username"],
    )
    return {"ok": True, "status": body.status, "plan": plan}


@app.delete("/api/attention/{attention_id}/plan")
def remove_plan(
    attention_id: str,
    request: Request,
    user=Depends(current_user),
    x_csrf_token: Optional[str] = Header(default=None),
):
    require_csrf(request, x_csrf_token)
    if user["role"] == "viewer":
        raise HTTPException(status_code=403, detail="El perfil Visualizador no puede eliminar planes.")
    rows = load_visible_rows(user)
    if not any(row["ID"] == attention_id for row in rows):
        raise HTTPException(status_code=404, detail="Evaluaci\u00f3n no encontrada.")
    if not delete_plan(attention_id, user["username"]):
        raise HTTPException(status_code=404, detail="No hay un plan activo para eliminar.")
    return {"ok": True}


@app.get("/api/admin/users")
def admin_users(user=Depends(current_user)):
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Acceso solo para administración.")
    return {
        "roles": [{"value": key, "label": value} for key, value in ROLES.items()],
        "assignable_roles": [
            {"value": key, "label": value}
            for key, value in ROLES.items()
            if key != "admin"
        ],
        "users": [dict(row) for row in list_users()],
    }


@app.post("/api/admin/users", status_code=status.HTTP_201_CREATED)
def create_user(
    body: UserBody,
    request: Request,
    user=Depends(current_user),
    x_csrf_token: Optional[str] = Header(default=None),
):
    require_csrf(request, x_csrf_token)
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Acceso solo para administración.")
    username = body.username.strip()
    if username.casefold() == FIXED_ADMIN_USERNAME or body.role == "admin":
        raise HTTPException(status_code=422, detail="El usuario admin es fijo y no se crea desde este formulario.")
    validate_user_assignment(username, body.role, (body.scope_name or "").strip(), body.password)
    try:
        password_hash = hash_password(body.password)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    try:
        save_user(
            username=username,
            full_name=body.full_name.strip(),
            password_hash=password_hash,
            role=body.role,
            scope_name=(body.scope_name or "").strip(),
        )
    except Exception as exc:
        logger.info("Could not create user: %s", exc)
        raise HTTPException(status_code=409, detail="Verifica que el usuario no exista.") from exc
    return {"ok": True}


@app.put("/api/admin/users/{username}")
def update_user(
    username: str,
    body: UserUpdateBody,
    request: Request,
    user=Depends(current_user),
    x_csrf_token: Optional[str] = Header(default=None),
):
    require_csrf(request, x_csrf_token)
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Acceso solo para administración.")
    target = username.strip()
    if not target:
        raise HTTPException(status_code=422, detail="Usuario no válido.")
    scope_name = (body.scope_name or "").strip()
    password = body.password if body.password else None
    validate_user_assignment(target, body.role, scope_name, password)
    try:
        password_hash = hash_password(password) if password else None
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    if not update_user_account(
        target,
        body.full_name.strip(),
        body.role,
        scope_name,
        password_hash,
    ):
        raise HTTPException(status_code=404, detail="El usuario ya no existe.")
    return {"ok": True}


@app.delete("/api/admin/users/{username}")
def remove_user(
    username: str,
    request: Request,
    user=Depends(current_user),
    x_csrf_token: Optional[str] = Header(default=None),
):
    require_csrf(request, x_csrf_token)
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Acceso solo para administración.")
    target = username.strip()
    if target.casefold() == FIXED_ADMIN_USERNAME:
        raise HTTPException(status_code=422, detail="El usuario admin fijo no se puede eliminar.")
    if not delete_user_account(target):
        raise HTTPException(status_code=404, detail="El usuario ya no existe.")
    return {"ok": True}


@app.post("/api/admin/users/import")
def import_users(
    body: UsersImportBody,
    request: Request,
    user=Depends(current_user),
    x_csrf_token: Optional[str] = Header(default=None),
):
    require_csrf(request, x_csrf_token)
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Acceso solo para administración.")
    users, skipped = parse_users_import(body.content)
    if not users and not skipped:
        raise HTTPException(status_code=422, detail="El archivo no contiene usuarios para importar.")
    created, duplicates = save_users_batch(users)
    skipped.extend(duplicates)
    return {
        "filename": body.filename,
        "created": len(created),
        "skipped": len(skipped),
        "created_users": created,
        "details": skipped[:100],
    }
